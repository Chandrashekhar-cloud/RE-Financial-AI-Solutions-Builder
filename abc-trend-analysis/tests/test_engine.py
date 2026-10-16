"""Pytest test suite for FinancialCalculationEngine.

Covers:
- Month ordering and reverse chronology handling
- Decimal arithmetic precision
- Subtotal handling (preventing double counting of S&M Payroll & Marketing)
- MoM calculations (absolute value denominator, division-by-zero suppression)
- TTM eligibility (11 vs 12 months, calendar continuity, 7 windows on 18 months)
- Accounting anomalies (Dec 2025 tax credit & SBC, non-profitability assertion)
- Negative line item reversals (Mar 2026)
- Interest misclassification in opex
- Data reconciliation checks
- Multi-company generalization
"""

import json
import pytest
from decimal import Decimal
from pathlib import Path
import openpyxl

from scripts.engine import FinancialCalculationEngine, to_cents
from scripts.models import Severity, AnomalyCategory


BASE_DIR = Path(__file__).resolve().parent.parent
ABC_WORKBOOK = BASE_DIR / "data" / "ABC Inc - Financial Statements.xlsx"
FIXTURES_DIR = BASE_DIR / "evals" / "fixtures"


class TestMonthOrdering:
    def test_reverse_order_detection_and_chronological_sorting(self):
        engine = FinancialCalculationEngine(str(ABC_WORKBOOK))
        result = engine.run_analysis()

        assert result.months_were_reversed is True, "Source file was reverse chronological"
        assert len(result.periods_analyzed) == 18

        # Oldest to newest
        assert result.periods_analyzed[0]["key"] == "2025-01"
        assert result.periods_analyzed[0]["label"] == "Jan 2025"
        assert result.periods_analyzed[-1]["key"] == "2026-06"
        assert result.periods_analyzed[-1]["label"] == "Jun 2026"

        # Strictly ascending order
        keys = [p["key"] for p in result.periods_analyzed]
        assert keys == sorted(keys), "Periods must be strictly ascending chronologically"

    def test_forward_order_workbook(self):
        engine = FinancialCalculationEngine(str(FIXTURES_DIR / "case1_normal_12m.xlsx"))
        result = engine.run_analysis()

        assert result.months_were_reversed is False
        assert result.periods_analyzed[0]["key"] == "2025-01"
        assert result.periods_analyzed[-1]["key"] == "2025-12"


class TestDecimalArithmeticPrecision:
    def test_to_cents_precision(self):
        assert to_cents(123.456) == Decimal("123.46")
        assert to_cents("1,234.50") == Decimal("1234.50")
        assert to_cents("($500.25)") == Decimal("-500.25")
        assert to_cents(None) == Decimal("0.00")
        assert to_cents("") == Decimal("0.00")

    def test_reconciliation_zero_floating_drift(self):
        engine = FinancialCalculationEngine(str(ABC_WORKBOOK))
        result = engine.run_analysis()

        for check in result.reconciliation_checks:
            assert check.is_valid is True, f"Reconciliation failed for {check.metric} in {check.period}"
            assert Decimal(check.difference) == Decimal("0.00"), f"Drift found: {check.difference}"


class TestSubtotalHandling:
    def test_no_double_counting_of_payroll_and_marketing(self):
        engine = FinancialCalculationEngine(str(ABC_WORKBOOK))
        result = engine.run_analysis()

        # For every month, verified that calc_total_opex == reported_total_opex
        opex_checks = [c for c in result.reconciliation_checks if c.metric == "Total Operating Expenses"]
        assert len(opex_checks) == 18
        assert all(c.is_valid for c in opex_checks)

        # In Jun 2026, reported opex is 468181.80
        jun_26_summary = next(s for s in result.monthly_summaries if s.period == "2026-06")
        assert jun_26_summary.total_opex == "468181.80"

        # Verify that subtotal components were identified
        sm_sub = next(li for li in result.line_items if li.account_name == "Total S & M Payroll")
        assert sm_sub.is_subtotal is True
        assert len(sm_sub.subtotal_components) > 0, "Subtotal components must be populated"

    def test_non_abc_synthetic_subtotals_and_reconciliation(self, tmp_path):
        import openpyxl
        fixture_path = tmp_path / "synthetic_non_abc_pnl.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Monthly P&L"

        ws.append(["Income Statement (Profit and Loss)", "", ""])
        ws.append(["Apex Dynamics Inc", "", ""])
        ws.append(["For period ended Feb 2025", "", ""])
        ws.append(["", "", ""])
        ws.append(["", "Account", "Jan 2025", "Feb 2025"])
        ws.append(["Income", "", "", ""])
        ws.append(["", "Enterprise Subscriptions", 200000, 220000])
        ws.append(["Total Income", "", 200000, 220000])
        ws.append(["Cost of Goods Sold", "", "", ""])
        ws.append(["", "Hosting & Infrastructure", 40000, 44000])
        ws.append(["Total Cost of Goods Sold", "", 40000, 44000])
        ws.append(["Gross Profit", "", 160000, 176000])
        ws.append(["Operating Expenses", "", "", ""])
        ws.append(["", "Sales Salaries", "", ""])
        ws.append(["", "AE Base Salaries", 30000, 32000])
        ws.append(["", "SDR Base Salaries", 20000, 21000])
        ws.append(["", "Total Sales Salaries", 50000, 53000])
        ws.append(["", "Growth & Advertising", "", ""])
        ws.append(["", "Search Ad Spend", 25000, 26000])
        ws.append(["", "Content Syndication", 15000, 16000])
        ws.append(["", "Total Growth & Advertising", 40000, 42000])
        ws.append(["", "Facilities Rent", 10000, 10000])
        ws.append(["Total Operating Expenses", "", 100000, 105000])
        ws.append(["Operating Income", "", 60000, 71000])
        ws.append(["Net Income", "", 60000, 71000])
        wb.save(str(fixture_path))

        engine = FinancialCalculationEngine(str(fixture_path))
        result = engine.run_analysis()

        # 1. Totals are NOT double-counted: reported total opex matches calculated total opex
        opex_checks = [c for c in result.reconciliation_checks if c.metric == "Total Operating Expenses"]
        assert len(opex_checks) == 2
        assert all(c.is_valid for c in opex_checks)
        assert result.monthly_summaries[0].total_opex == "100000.00"
        assert result.monthly_summaries[1].total_opex == "105000.00"

        # 2. Subtotal components identified correctly
        sales_sub = next(li for li in result.line_items if li.account_name == "Total Sales Salaries")
        growth_sub = next(li for li in result.line_items if li.account_name == "Total Growth & Advertising")
        assert sales_sub.is_subtotal is True and len(sales_sub.subtotal_components) == 2
        assert growth_sub.is_subtotal is True and len(growth_sub.subtotal_components) == 2

        # 3. Subtotal reconciliation checks exist and pass
        metrics = {c.metric for c in result.reconciliation_checks}
        assert "Total Sales Salaries" in metrics, f"Missing Total Sales Salaries check, got: {metrics}"
        assert "Total Growth & Advertising" in metrics, f"Missing Total Growth & Advertising check, got: {metrics}"
        assert all(c.is_valid for c in result.reconciliation_checks)


class TestMoMCalculations:
    def test_mom_suppressed_for_first_month(self):
        engine = FinancialCalculationEngine(str(ABC_WORKBOOK))
        result = engine.run_analysis()

        first = result.monthly_summaries[0]
        assert first.period == "2025-01"
        assert first.revenue_mom_diff is None
        assert first.revenue_mom_pct is None

    def test_mom_percentage_formula_with_positive_prior(self):
        engine = FinancialCalculationEngine(str(ABC_WORKBOOK))
        result = engine.run_analysis()

        # Jan 2025: 167587.50, Feb 2025: 176651.19
        feb = result.monthly_summaries[1]
        assert feb.period == "2025-02"
        assert feb.revenue_mom_diff == "9063.69"
        assert feb.revenue_mom_pct == "5.41%"

    def test_mom_percentage_with_negative_prior(self):
        engine = FinancialCalculationEngine(str(ABC_WORKBOOK))
        result = engine.run_analysis()

        # Dec 2025 Total Opex: -402142.09
        # Jan 2026 Total Opex: 487711.23
        # Diff = 487711.23 - (-402142.09) = +889853.32
        # Pct = (889853.32 / abs(-402142.09)) * 100 = 221.28%
        jan_26 = next(s for s in result.monthly_summaries if s.period == "2026-01")
        assert jan_26.opex_mom_diff == "889853.32"
        assert jan_26.opex_mom_pct == "221.28%"


class TestTTMEligibility:
    def test_exactly_seven_ttm_windows_on_18_months(self):
        engine = FinancialCalculationEngine(str(ABC_WORKBOOK))
        result = engine.run_analysis()

        valid_windows = [w for w in result.ttm_windows if w.is_valid]
        assert len(valid_windows) == 7, "Exactly 7 TTM windows must be generated for 18 months"

        expected_end_periods = [
            "2025-12", "2026-01", "2026-02", "2026-03", "2026-04", "2026-05", "2026-06"
        ]
        actual_end_periods = [w.end_period for w in valid_windows]
        assert actual_end_periods == expected_end_periods

        # Each window must contain exactly 12 months
        for w in valid_windows:
            assert len(w.months_included) == 12

    def test_11_months_dataset_ineligible_for_ttm(self):
        engine = FinancialCalculationEngine(str(FIXTURES_DIR / "case4_11_months.xlsx"))
        result = engine.run_analysis()

        valid_windows = [w for w in result.ttm_windows if w.is_valid]
        assert len(valid_windows) == 0, "No TTM windows should be generated for 11 months"
        assert any(a.category == AnomalyCategory.LIMITATION for a in result.anomalies)

    def test_missing_month_gap_invalidates_ttm_across_gap(self):
        engine = FinancialCalculationEngine(str(FIXTURES_DIR / "case6_missing_month.xlsx"))
        result = engine.run_analysis()

        assert len(result.missing_period_gaps) == 1
        assert result.missing_period_gaps[0]["missing_month"] == "2025-03"

        valid_windows = [w for w in result.ttm_windows if w.is_valid]
        assert len(valid_windows) == 0


class TestAccountingAnomalies:
    def test_december_2025_tax_benefit_and_sbc_flags(self):
        engine = FinancialCalculationEngine(str(ABC_WORKBOOK))
        result = engine.run_analysis()

        # Check critical tax benefit flag
        tax_flag = next(
            (a for a in result.anomalies if a.severity == Severity.CRITICAL and "tax benefit" in a.account.lower()),
            None
        )
        assert tax_flag is not None
        assert tax_flag.period == "2025-12"
        assert "-1244429.84" in str(tax_flag.amount)
        assert "DID NOT become genuinely operationally profitable" in tax_flag.impact

        # Check SBC flag
        sbc_flag = next(
            (a for a in result.anomalies if "stock-based" in a.account.lower()),
            None
        )
        assert sbc_flag is not None
        assert sbc_flag.period == "2025-12"
        assert "334660.53" in str(sbc_flag.amount)

    def test_december_2025_normalized_profitability_rejection(self):
        engine = FinancialCalculationEngine(str(ABC_WORKBOOK))
        result = engine.run_analysis()

        dec = next(s for s in result.monthly_summaries if s.period == "2025-12")
        reported_oi = Decimal(dec.operating_income)
        normalized_oi = Decimal(dec.normalized_operating_income)

        # Reported shows +$534,551.41
        assert reported_oi > Decimal("0.00")
        assert reported_oi == Decimal("534551.41")

        # Normalized shows a core operating loss
        assert normalized_oi < Decimal("0.00")
        assert normalized_oi == Decimal("-381241.27"), "Normalized core operating loss must equal verified -$381,241.27"

    def test_march_2026_reversals(self):
        engine = FinancialCalculationEngine(str(ABC_WORKBOOK))
        result = engine.run_analysis()

        reversals = [a for a in result.anomalies if a.category == AnomalyCategory.REVERSAL and a.period == "2026-03"]
        assert len(reversals) >= 2, "Must flag March 2026 negative reversals"

        accounts = [r.account for r in reversals]
        assert "Partnership Income" in accounts
        assert "White-Label" in accounts

    def test_interest_in_opex_flagged(self):
        engine = FinancialCalculationEngine(str(ABC_WORKBOOK))
        result = engine.run_analysis()

        misclass = [a for a in result.anomalies if a.category == AnomalyCategory.MISCLASSIFICATION]
        accounts = [m.account for m in misclass]
        assert "Interest Income" in accounts
        assert "Interest Expense" in accounts


class TestGeneralizationAndIntegrity:
    def test_different_company_chart_of_accounts(self):
        engine = FinancialCalculationEngine(str(FIXTURES_DIR / "case8_different_company.xlsx"))
        result = engine.run_analysis()

        assert len(result.periods_analyzed) == 12
        assert len([w for w in result.ttm_windows if w.is_valid]) == 1
        assert all(rc.is_valid for rc in result.reconciliation_checks)

    def test_reconciliation_failure_detection(self):
        engine = FinancialCalculationEngine(str(FIXTURES_DIR / "case7_inconsistent_totals.xlsx"))
        result = engine.run_analysis()

        failed_checks = [rc for rc in result.reconciliation_checks if not rc.is_valid]
        assert len(failed_checks) > 0, "Corrupted total must trigger a failed reconciliation check"
        assert any(a.category == AnomalyCategory.RECONCILIATION_DISCREPANCY for a in result.anomalies)

    def test_unmapped_accounts_reporting(self):
        engine = FinancialCalculationEngine(str(ABC_WORKBOOK))
        result = engine.run_analysis()

        # In ABC Inc, 100% of line items match standard chart of accounts
        assert len(result.unmapped_accounts) == 0, "ABC Inc has zero unmapped accounts"

        # Explicit classification test
        assert engine._classify_department("Mysterious Wire Transfer", "Operating Expenses", False, False) == "Unmapped"
        assert engine._classify_department("Payroll Service Fee", "Operating Expenses", False, False) == "General & Administrative"

    def test_unmapped_opex_included_in_normalized_core_opex(self):
        """Verifies that unmapped operating expenses are retained in normalized core opex rather than silently excluded."""
        engine = FinancialCalculationEngine(str(FIXTURES_DIR / "case8_different_company.xlsx"))
        result = engine.run_analysis()

        # 1. Unmapped accounts remain flagged for CPA review
        assert "Engineering Salaries" in result.unmapped_accounts
        assert "Growth & Customer Acquisition" in result.unmapped_accounts
        assert len(result.unmapped_accounts) == 2

        # 2. Total Opex in Jan 2025 is $170,000 (80k engineering + 60k growth + 30k exec)
        jan_summary = result.monthly_summaries[0]
        assert jan_summary.total_opex == "170000.00"

        # 3. Normalized core opex must retain the unmapped $140,000 expenses (matching total opex of $170,000.00)
        assert jan_summary.normalized_core_opex == "170000.00"
        assert jan_summary.normalized_operating_income == "70000.00"

        # 4. Latest TTM normalized core operating income equals reported operating income ($1,533,000.00)
        latest_ttm = next(w for w in result.ttm_windows if w.is_valid)
        assert latest_ttm.total_opex == "2040000.00"
        assert latest_ttm.normalized_core_opex == "2040000.00"
        assert latest_ttm.operating_income == "1533000.00"
        assert latest_ttm.normalized_operating_income == "1533000.00"

    def test_tax_benefit_anomaly_computes_dynamic_amounts_for_non_abc_fixture(self, tmp_path):
        """Synthetic non-ABC fixture triggers tax benefit anomaly with its own computed amounts."""
        src_path = FIXTURES_DIR / "case8_different_company.xlsx"
        wb = openpyxl.load_workbook(src_path)
        ws = wb.active

        ws.insert_rows(20)
        ws.cell(20, 2, "Income Tax Benefit")

        for c in range(3, 15):
            ws.cell(20, c, 0.0)

        # For Nov 2025 (col 13):
        # Income = 100,000
        ws.cell(8, 13, 100000.0)
        ws.cell(9, 13, 0.0)
        ws.cell(10, 13, 100000.0)
        # COGS = 20,000
        ws.cell(12, 13, 20000.0)
        ws.cell(13, 13, 0.0)
        ws.cell(14, 13, 20000.0)
        # GP = 80,000
        ws.cell(15, 13, 80000.0)
        # Core Opex: R17 = 0, R18 = 0, R19 (Executive & Administrative) = 150,000
        ws.cell(17, 13, 0.0)
        ws.cell(18, 13, 0.0)
        ws.cell(19, 13, 150000.0)
        # Tax Benefit (R20) = -250,000
        ws.cell(20, 13, -250000.0)
        # Total Opex (R21) = 150,000 - 250,000 = -100,000
        ws.cell(21, 13, -100000.0)
        # Operating Income (R22) = GP (80,000) - Total Opex (-100,000) = 180,000
        ws.cell(22, 13, 180000.0)
        # Net Income (R23) = 180,000
        ws.cell(23, 13, 180000.0)

        fixture_file = tmp_path / "test_corp_tax_anomaly.xlsx"
        wb.save(fixture_file)

        engine = FinancialCalculationEngine(str(fixture_file))
        result = engine.run_analysis()

        tax_anomalies = [a for a in result.anomalies if "tax benefit" in a.account.lower()]
        assert len(tax_anomalies) == 1, "Must detect exactly 1 tax benefit anomaly"
        anomaly = tax_anomalies[0]

        # 1. Assert anomaly text contains the fixture's own computed amounts
        assert "$250,000.00" in anomaly.description
        assert "-$100,000.00" in anomaly.description
        assert "+$180,000.00" in anomaly.description
        assert "Nov 2025" in anomaly.impact
        assert "-$70,000.00" in anomaly.impact

        # 2. Assert anomaly text does NOT contain any ABC figures
        abc_figures = [
            "-402,142.09",
            "534,551.41",
            "-381,241.27",
            "-375,217.90",
            "1,244,429.84",
            "Dec 2025",
        ]
        for figure in abc_figures:
            assert figure not in anomaly.description, f"ABC figure {figure} leaked into description"
            assert figure not in anomaly.impact, f"ABC figure {figure} leaked into impact"

    def test_tax_benefit_threshold_override(self, tmp_path):
        """A synthetic non-ABC fixture with a $15,000 tax credit.

        Assert NO anomaly is raised under default thresholds (100,000.00),
        and assert the anomaly IS raised when thresholds are overridden to 10,000.00.
        """
        src_path = FIXTURES_DIR / "case8_different_company.xlsx"
        wb = openpyxl.load_workbook(src_path)
        ws = wb.active

        ws.insert_rows(20)
        ws.cell(20, 2, "Income Tax Benefit")

        for c in range(3, 15):
            ws.cell(20, c, 0.0)

        # For Nov 2025 (col 13), add a $15,000 tax credit (-15,000.00)
        # Income = 100,000
        ws.cell(8, 13, 100000.0)
        ws.cell(9, 13, 0.0)
        ws.cell(10, 13, 100000.0)
        # COGS = 20,000
        ws.cell(12, 13, 20000.0)
        ws.cell(13, 13, 0.0)
        ws.cell(14, 13, 20000.0)
        # GP = 80,000
        ws.cell(15, 13, 80000.0)
        # Core Opex: R17 = 0, R18 = 0, R19 (Executive & Administrative) = 50,000
        ws.cell(17, 13, 0.0)
        ws.cell(18, 13, 0.0)
        ws.cell(19, 13, 50000.0)
        # Tax Benefit = -15,000.00
        ws.cell(20, 13, -15000.0)
        # Total Opex (R21) = 50,000 - 15,000 = 35,000
        ws.cell(21, 13, 35000.0)
        # Operating Income (R22) = GP (80,000) - Total Opex (35,000) = 45,000
        ws.cell(22, 13, 45000.0)
        # Net Income (R23) = 45,000
        ws.cell(23, 13, 45000.0)

        fixture_file = tmp_path / "test_corp_15k_tax.xlsx"
        wb.save(fixture_file)

        # 1. Under default thresholds (100,000 credit threshold): NO anomaly
        default_engine = FinancialCalculationEngine(str(fixture_file))
        default_result = default_engine.run_analysis()
        tax_anomalies_default = [a for a in default_result.anomalies if "tax benefit" in a.account.lower()]
        assert len(tax_anomalies_default) == 0, "No tax benefit anomaly should trigger under default 100k threshold"

        # 2. Under overridden thresholds (10,000 credit threshold): Anomaly IS raised
        override_config = tmp_path / "override_thresholds.json"
        override_config.write_text(json.dumps({
            "tax_benefit_credit": 10000.00,
            "stock_comp_spike": 50000.00,
            "expense_reversal": 500.00
        }), encoding="utf-8")

        override_engine = FinancialCalculationEngine(str(fixture_file), thresholds_path=override_config)
        override_result = override_engine.run_analysis()
        tax_anomalies_override = [a for a in override_result.anomalies if "tax benefit" in a.account.lower()]
        assert len(tax_anomalies_override) == 1, "Tax benefit anomaly must trigger under overridden 10k threshold"
        assert "$15,000.00" in tax_anomalies_override[0].description

    def test_missing_or_malformed_thresholds_raises_clear_error(self, tmp_path):
        """A missing or malformed thresholds file must fail with a clear error, not fall back silently."""
        src_path = FIXTURES_DIR / "case8_different_company.xlsx"

        # Missing file
        missing_file = tmp_path / "non_existent_thresholds.json"
        with pytest.raises(FileNotFoundError, match="Anomaly thresholds configuration file not found"):
            FinancialCalculationEngine(str(src_path), thresholds_path=missing_file)

        # Malformed JSON (syntax error)
        malformed_json = tmp_path / "malformed.json"
        malformed_json.write_text("{invalid_json: true", encoding="utf-8")
        with pytest.raises(ValueError, match="Malformed JSON in anomaly thresholds configuration file"):
            FinancialCalculationEngine(str(src_path), thresholds_path=malformed_json)

        # Missing required key
        missing_key_json = tmp_path / "missing_key.json"
        missing_key_json.write_text(json.dumps({
            "tax_benefit_credit": 100000.00,
            "expense_reversal": 500.00
        }), encoding="utf-8")
        with pytest.raises(ValueError, match="Missing required threshold key 'stock_comp_spike'"):
            FinancialCalculationEngine(str(src_path), thresholds_path=missing_key_json)

        # Non-numeric value
        invalid_val_json = tmp_path / "invalid_val.json"
        invalid_val_json.write_text(json.dumps({
            "tax_benefit_credit": "not-a-number",
            "stock_comp_spike": 50000.00,
            "expense_reversal": 500.00
        }), encoding="utf-8")
        with pytest.raises(ValueError, match="Invalid numeric threshold value for 'tax_benefit_credit'"):
            FinancialCalculationEngine(str(src_path), thresholds_path=invalid_val_json)


@pytest.fixture(scope="module")
def real_abc_analysis():
    engine = FinancialCalculationEngine(str(ABC_WORKBOOK))
    return engine.run_analysis()


class TestRealWorkbookRegression:
    """End-to-end regression test for the supplied real ABC Inc financial statements workbook.

    Verifies independently specified expected values (period coverage, Dec 2025 normalized loss,
    TTM revenue and gross margins, Mar 2026 reversals, and reconciliation integrity) without
    deriving them dynamically from the engine under test.
    """

    def test_reporting_period_coverage(self, real_abc_analysis):
        """Verifies full 18-month coverage (Jan 2025 - Jun 2026) and chronological normalization."""
        # Expected periods independently specified from workbook dates
        expected_periods = [
            "2025-01", "2025-02", "2025-03", "2025-04", "2025-05", "2025-06",
            "2025-07", "2025-08", "2025-09", "2025-10", "2025-11", "2025-12",
            "2026-01", "2026-02", "2026-03", "2026-04", "2026-05", "2026-06",
        ]
        actual_periods = [p["key"] for p in real_abc_analysis.periods_analyzed]
        assert actual_periods == expected_periods
        assert len(real_abc_analysis.periods_analyzed) == 18
        assert real_abc_analysis.months_were_reversed is True
        assert real_abc_analysis.periods_analyzed[0]["key"] == "2025-01"
        assert real_abc_analysis.periods_analyzed[-1]["key"] == "2026-06"

    def test_reconciliation_integrity_and_unmapped(self, real_abc_analysis):
        """Verifies 100% mathematical reconciliation across 126 checks and 0 unmapped accounts."""
        assert len(real_abc_analysis.reconciliation_checks) == 126
        assert all(rc.is_valid for rc in real_abc_analysis.reconciliation_checks)
        assert real_abc_analysis.unmapped_accounts == []

    def test_december_2025_independently_specified_values(self, real_abc_analysis):
        """Verifies independently specified financial values for December 2025."""
        dec_summary = next(s for s in real_abc_analysis.monthly_summaries if s.period == "2025-12")

        # Independently specified expected values from source accounting statements:
        assert dec_summary.revenue == "190029.49"
        assert dec_summary.cogs == "57620.17"
        assert dec_summary.gross_profit == "132409.32"
        assert dec_summary.total_opex == "-402142.09"
        assert dec_summary.operating_income == "534551.41"
        assert dec_summary.normalized_core_opex == "513650.59"
        assert dec_summary.normalized_operating_income == "-381241.27"

        # Verify false profitability rejection: reported > 0, but normalized core < 0
        assert Decimal(dec_summary.operating_income) > Decimal("0.00")
        assert Decimal(dec_summary.normalized_operating_income) < Decimal("0.00")

        # Verify key anomaly amounts
        tax_anomalies = [
            a for a in real_abc_analysis.anomalies
            if a.account == "Income Tax Benefit" and a.period == "2025-12"
        ]
        assert len(tax_anomalies) == 1
        assert tax_anomalies[0].amount == "-1244429.84"
        assert tax_anomalies[0].severity == Severity.CRITICAL
        assert tax_anomalies[0].category == AnomalyCategory.ACCOUNTING_ANOMALY

        sbc_anomalies = [
            a for a in real_abc_analysis.anomalies
            if a.account == "Stock-Based Compensation Expense" and a.period == "2025-12"
        ]
        assert len(sbc_anomalies) == 1
        assert sbc_anomalies[0].amount == "334660.53"
        assert sbc_anomalies[0].severity == Severity.WARNING
        assert sbc_anomalies[0].category == AnomalyCategory.ACCOUNTING_ANOMALY

    def test_march_2026_reversals_independently_specified(self, real_abc_analysis):
        """Verifies independently specified credit memo reversals and interest misclassification in March 2026."""
        mar_anomalies = [a for a in real_abc_analysis.anomalies if a.period == "2026-03"]
        reversal_map = {a.account: a.amount for a in mar_anomalies if a.category == AnomalyCategory.REVERSAL}

        assert reversal_map["Partnership Income"] == "-1608.09"
        assert reversal_map["White-Label"] == "-550.00"
        assert reversal_map["Sales Commission"] == "-3629.87"

        # Interest misclassification across statement
        misclass = [a for a in real_abc_analysis.anomalies if a.category == AnomalyCategory.MISCLASSIFICATION]
        misclass_accounts = [m.account for m in misclass]
        assert "Interest Income" in misclass_accounts
        assert "Interest Expense" in misclass_accounts

    def test_trailing_twelve_months_latest_window(self, real_abc_analysis):
        """Verifies independently calculated TTM metrics for the latest 12-month window (ending Jun 2026)."""
        valid_windows = [w for w in real_abc_analysis.ttm_windows if w.is_valid]
        assert len(valid_windows) == 7

        latest_ttm = next(w for w in valid_windows if w.end_period == "2026-06")
        assert latest_ttm.start_period == "2025-07"
        assert len(latest_ttm.months_included) == 12

        # Independently calculated sums across Jul 2025 - Jun 2026:
        assert latest_ttm.revenue == "2303975.51"
        assert latest_ttm.cogs == "703547.60"
        assert latest_ttm.gross_profit == "1600427.91"
        assert latest_ttm.gross_margin_pct == "69.46%"
        assert latest_ttm.operating_income == "-3089779.59"
        assert latest_ttm.normalized_operating_income == "-4064377.01"

