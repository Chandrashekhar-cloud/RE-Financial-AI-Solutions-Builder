"""
QuickBooks Online Profit and Loss Report Adapter.

Normalizes Intuit QuickBooks Online ProfitAndLoss API JSON responses into
the canonical financial statement structure consumed by FinancialCalculationEngine.
Enables end-to-end trend analysis on live or exported QuickBooks data without
modifying core calculation or reconciliation logic.
"""

import json
import tempfile
from pathlib import Path
from typing import Dict, Any, List, Tuple, Optional, Union
import openpyxl

from scripts.engine import FinancialCalculationEngine
from scripts.models import FinancialAnalysisResult
from scripts.reporter import FinancialReportGenerator


class QBOReportAdapter:
    """Adapter that normalizes QuickBooks Online Profit and Loss reports into

    the canonical schema expected by FinancialCalculationEngine.
    """

    def __init__(
        self,
        report_source: Union[Dict[str, Any], str, Path],
        thresholds_path: Optional[Union[str, Path]] = None,
    ):
        """Initialize adapter with QBO report data.

        Args:
            report_source: Either a dictionary representing the QBO report JSON,
                           a JSON string, or a path to a JSON file.
            thresholds_path: Optional path to custom anomaly thresholds JSON configuration file.
        """
        self.thresholds_path = thresholds_path
        if isinstance(report_source, (str, Path)) and Path(str(report_source)).exists():
            with open(report_source, "r", encoding="utf-8") as f:
                self.raw_report: Dict[str, Any] = json.load(f)
        elif isinstance(report_source, str):
            self.raw_report = json.loads(report_source)
        elif isinstance(report_source, dict):
            self.raw_report = report_source
        else:
            raise ValueError(f"Unsupported report source type: {type(report_source)}")

        self._unwrap_mcp_envelope_if_present()
        self._validate_report_structure()
        self.header = self.raw_report.get("Header", {})
        self.company_name = self._extract_company_name()
        self.statement_name = self.header.get("ReportName", "Income Statement (Profit and Loss)")
        self.start_period = self.header.get("StartPeriod", "")
        self.end_period = self.header.get("EndPeriod", "")
        self.month_columns: List[Tuple[int, str]] = self._extract_month_columns()

    def _unwrap_mcp_envelope_if_present(self) -> None:
        """Unwrap MCP tool response envelopes if the input was captured directly from an MCP tool call."""
        if not isinstance(self.raw_report, dict):
            return

        # 1. MCP tool handler format: { "result": { "Header": ..., "Columns": ..., "Rows": ... }, "isError": false }
        if "result" in self.raw_report and isinstance(self.raw_report["result"], dict) and "Columns" in self.raw_report["result"]:
            self.raw_report = self.raw_report["result"]
            return

        # 2. MCP client/Claude response envelope: { "content": [{ "type": "text", "text": "..." }] }
        if "content" in self.raw_report and isinstance(self.raw_report["content"], list):
            for item in self.raw_report["content"]:
                if isinstance(item, dict) and item.get("type") == "text":
                    text = item.get("text", "").strip()
                    if text.startswith("{") and ("Columns" in text or "Rows" in text):
                        try:
                            parsed = json.loads(text)
                            if isinstance(parsed, dict) and ("Columns" in parsed or "Rows" in parsed):
                                self.raw_report = parsed
                                return
                        except Exception:
                            pass

    def _validate_report_structure(self) -> None:
        """Verify report has minimum expected Intuit V3 structure."""
        if not isinstance(self.raw_report, dict):
            raise ValueError("QBO report must be a dictionary object.")
        if "Columns" not in self.raw_report or "Rows" not in self.raw_report:
            raise ValueError("Invalid QBO report: missing 'Columns' or 'Rows' root objects.")

    def _extract_company_name(self) -> str:
        """Extract company name or fallback to ABC Inc."""
        # Intuit sometimes includes CompanyName in Header or Options
        if "CompanyName" in self.header:
            return str(self.header["CompanyName"]).strip()
        options = self.header.get("Option", [])
        for opt in options:
            if opt.get("Name") == "CompanyName":
                return str(opt.get("Value", "")).strip()
        return "ABC Inc"

    def _extract_month_columns(self) -> List[Tuple[int, str]]:
        """Identify monthly data columns from Columns.Column definition.

        Returns list of (column_index_in_qbo, month_title).
        Excludes Account column (index 0) and Total column.
        """
        col_defs = self.raw_report.get("Columns", {}).get("Column", [])
        extracted = []
        for idx, col in enumerate(col_defs):
            title = col.get("ColTitle", "").strip()
            col_type = col.get("ColType", "")
            # Skip account title column and aggregate summary column
            if not title or title.lower() == "total":
                continue
            if col_type == "Money" or title:
                extracted.append((idx, title))
        return extracted

    def to_workbook(self, output_path: Optional[Union[str, Path]] = None) -> Path:
        """Normalize QBO P&L report into a canonical openpyxl Excel workbook.

        Args:
            output_path: Optional destination path. If None, a temporary file is created.

        Returns:
            Path to the written Excel workbook.
        """
        if output_path:
            out_file = Path(output_path)
            out_file.parent.mkdir(parents=True, exist_ok=True)
        else:
            fd, tmp_name = tempfile.mkstemp(prefix="qbo_norm_", suffix=".xlsx")
            import os
            os.close(fd)
            out_file = Path(tmp_name)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Income Statement"

        # Row 1-3 Metadata
        ws.cell(1, 1, "Income Statement (Profit and Loss)")
        ws.cell(2, 1, self.company_name)
        as_of = f"For the period ended {self.end_period}" if self.end_period else "Financial Statements"
        ws.cell(3, 1, as_of)

        # Row 5 Headers
        ws.cell(5, 2, "Account")
        for target_c, (_, m_title) in enumerate(self.month_columns, start=3):
            ws.cell(5, target_c, m_title)

        cur_row = 7

        def write_values(target_r: int, col_data_list: List[Dict[str, Any]]):
            """Write monthly numbers across target row."""
            for target_c, (q_idx, _) in enumerate(self.month_columns, start=3):
                val_str = "0.00"
                if q_idx < len(col_data_list):
                    val_str = col_data_list[q_idx].get("value", "0.00")
                try:
                    num_val = float(val_str)
                except (ValueError, TypeError):
                    num_val = 0.0
                ws.cell(target_r, target_c, num_val)

        def walk_section(sec_dict: Dict[str, Any]):
            nonlocal cur_row
            raw_hdr = sec_dict.get("Header", {}).get("ColData", [{}])[0].get("value", "").strip()
            # Canonical section name mapping
            if raw_hdr.lower() == "expenses":
                section_title = "Operating Expenses"
            else:
                section_title = raw_hdr

            if section_title:
                ws.cell(cur_row, 1, section_title)
                cur_row += 1

            sub_rows = sec_dict.get("Rows", {}).get("Row", [])
            for sr in sub_rows:
                stype = sr.get("type")
                if stype == "Section":
                    # Nested section (e.g. S & M Payroll or Marketing)
                    sub_hdr = sr.get("Header", {}).get("ColData", [{}])[0].get("value", "").strip()
                    if sub_hdr:
                        ws.cell(cur_row, 2, sub_hdr)
                        cur_row += 1

                    for child_item in sr.get("Rows", {}).get("Row", []):
                        acc_name = child_item.get("ColData", [{}])[0].get("value", "").strip()
                        ws.cell(cur_row, 2, acc_name)
                        write_values(cur_row, child_item.get("ColData", []))
                        cur_row += 1

                    sub_sum = sr.get("Summary", {}).get("ColData", [])
                    if sub_sum:
                        sum_title = sub_sum[0].get("value", "").strip()
                        ws.cell(cur_row, 2, sum_title)
                        write_values(cur_row, sub_sum)
                        cur_row += 1

                elif stype == "Data" or "ColData" in sr:
                    acc_name = sr.get("ColData", [{}])[0].get("value", "").strip()
                    ws.cell(cur_row, 2, acc_name)
                    write_values(cur_row, sr.get("ColData", []))
                    cur_row += 1

            # Section summary
            sec_sum = sec_dict.get("Summary", {}).get("ColData", [])
            if sec_sum:
                sum_title = sec_sum[0].get("value", "").strip()
                if sum_title.lower() == "total expenses":
                    sum_title = "Total Operating Expenses"
                elif sum_title.lower() == "net operating income":
                    sum_title = "Operating Income"
                ws.cell(cur_row, 1, sum_title)
                write_values(cur_row, sec_sum)
                cur_row += 2

        # Traverse all top-level rows
        top_rows = self.raw_report.get("Rows", {}).get("Row", [])
        for top_item in top_rows:
            if top_item.get("type") == "Section":
                walk_section(top_item)

        wb.save(out_file)
        return out_file

    def run_analysis(
        self,
        output_workbook_path: Optional[Union[str, Path]] = None,
        thresholds_path: Optional[Union[str, Path]] = None,
    ) -> FinancialAnalysisResult:
        """Convert QBO report to normalized workbook and run deterministic engine.

        Args:
            output_workbook_path: Optional path to save intermediate normalized workbook.
            thresholds_path: Optional override path for anomaly detection thresholds JSON.

        Returns:
            FinancialAnalysisResult containing all summaries, TTMs, and reconciliations.
        """
        wb_path = self.to_workbook(output_workbook_path)
        thresh = thresholds_path or self.thresholds_path
        engine = FinancialCalculationEngine(str(wb_path), source_type="quickbooks", thresholds_path=thresh)
        return engine.run_analysis()

    def to_json(
        self,
        output_json_path: Optional[Union[str, Path]] = None,
        output_report_path: Optional[Union[str, Path]] = None,
        output_workbook_path: Optional[Union[str, Path]] = None,
    ) -> Dict[str, Any]:
        """Execute full analysis and optionally persist analysis.json and report.md.

        Returns:
            Dictionary matching the analysis.json contract.
        """
        result = self.run_analysis(output_workbook_path)
        result_dict = result.to_dict()

        if output_json_path:
            json_file = Path(output_json_path)
            json_file.parent.mkdir(parents=True, exist_ok=True)
            with open(json_file, "w", encoding="utf-8") as f:
                json.dump(result_dict, f, indent=2)

        if output_report_path:
            rep_file = Path(output_report_path)
            rep_file.parent.mkdir(parents=True, exist_ok=True)
            generator = FinancialReportGenerator(result)
            report_content = generator.generate_markdown()
            with open(rep_file, "w", encoding="utf-8") as f:
                f.write(report_content)

        return result_dict
